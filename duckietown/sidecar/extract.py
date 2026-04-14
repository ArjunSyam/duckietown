"""
extract.py — Multimodal content extraction.

For images: returns the file_path so ingest.py can caption + embed via Jina vision.
For text files: extracts and chunks text content.
"""

import csv
import mimetypes
import os
from pathlib import Path
from typing import Any, Dict, List

CHUNK_SIZE = 500
CHUNK_OVERLAP = 100


def _chunk_text(text: str) -> List[str]:
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + CHUNK_SIZE
        chunks.append(text[start:end])
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks


def extract_pdf(path: str) -> List[Dict[str, Any]]:
    import pdfplumber

    chunks = []
    with pdfplumber.open(path) as pdf:
        for page_num, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            for chunk in _chunk_text(text):
                chunks.append(
                    {
                        "text": chunk,
                        "chunk_index": len(chunks),
                        "modality": "text",
                        "page": page_num + 1,
                    }
                )
    return chunks


def extract_docx(path: str) -> List[Dict[str, Any]]:
    from docx import Document

    doc = Document(path)
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [
        {"text": c, "chunk_index": i, "modality": "text", "page": None}
        for i, c in enumerate(_chunk_text(full_text))
    ]


def extract_xlsx(path: str) -> List[Dict[str, Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                rows.append(row_text)
    full_text = "\n".join(rows)
    return [
        {"text": c, "chunk_index": i, "modality": "text", "page": None}
        for i, c in enumerate(_chunk_text(full_text))
    ]


def extract_csv(path: str) -> List[Dict[str, Any]]:
    rows = []
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(" | ".join(row))
    full_text = "\n".join(rows)
    return [
        {"text": c, "chunk_index": i, "modality": "text", "page": None}
        for i, c in enumerate(_chunk_text(full_text))
    ]


def extract_text(path: str) -> List[Dict[str, Any]]:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    return [
        {"text": c, "chunk_index": i, "modality": "text", "page": None}
        for i, c in enumerate(_chunk_text(content))
    ]


def extract_image(path: str) -> List[Dict[str, Any]]:
    """
    Return the file path for ingest.py to handle captioning + Jina vision embedding.
    The 'text' stub is only used as fallback if captioning fails.
    """
    name = Path(path).stem.replace("_", " ").replace("-", " ")
    return [
        {
            "text": f"Image: {name}",
            "chunk_index": 0,
            "modality": "image",
            "page": None,
            "file_path": path,
        }
    ]


def extract_audio(path: str) -> List[Dict[str, Any]]:
    name = Path(path).stem.replace("_", " ").replace("-", " ")
    return [
        {
            "text": f"Audio recording: {name}",
            "chunk_index": 0,
            "modality": "audio",
            "page": None,
        }
    ]


def extract_video(path: str) -> List[Dict[str, Any]]:
    name = Path(path).stem.replace("_", " ").replace("-", " ")
    return [
        {
            "text": f"Video: {name}",
            "chunk_index": 0,
            "modality": "video",
            "page": None,
        }
    ]


def extract(path: str) -> List[Dict[str, Any]]:
    mime, _ = mimetypes.guess_type(path)
    mime = mime or ""
    ext = Path(path).suffix.lower()

    try:
        if mime == "application/pdf" or ext == ".pdf":
            return extract_pdf(path)
        elif mime in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ) or ext in (".docx", ".doc"):
            return extract_docx(path)
        elif mime in (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        ) or ext in (".xlsx", ".xls"):
            return extract_xlsx(path)
        elif ext == ".csv" or mime == "text/csv":
            return extract_csv(path)
        elif (mime and mime.startswith("text/")) or ext in (
            ".txt",
            ".md",
            ".py",
            ".js",
            ".ts",
            ".json",
            ".yaml",
            ".yml",
        ):
            return extract_text(path)
        elif mime and mime.startswith("image/"):
            return extract_image(path)
        elif mime and mime.startswith("audio/"):
            return extract_audio(path)
        elif mime and mime.startswith("video/"):
            return extract_video(path)
        else:
            name = Path(path).name
            return [
                {
                    "text": f"File: {name}",
                    "chunk_index": 0,
                    "modality": "other",
                    "page": None,
                }
            ]
    except Exception as e:
        print(f"[extract] Error on {path}: {e}")
        name = Path(path).name
        return [
            {
                "text": f"File: {name}",
                "chunk_index": 0,
                "modality": "other",
                "page": None,
            }
        ]
